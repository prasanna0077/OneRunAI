"""Write branded, analyst-friendly Excel workbooks (README + data sheets + data dictionary)."""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="00AEEF")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
BODY_FONT = Font(name="Arial", size=10)

# column -> (description, unit)
DICTIONARY = {
    "bank_id": ("Unique institution id", "-"), "bank_name": ("Institution name (fictional)", "-"),
    "bank_type": ("Institution category", "-"), "region": ("Primary operating region", "-"),
    "total_assets_cr": ("Total assets", "INR Cr"), "deposits_cr": ("Total deposits", "INR Cr"),
    "advances_cr": ("Gross advances", "INR Cr"), "networth_cr": ("Net worth", "INR Cr"),
    "roa_pct": ("Return on assets", "%"), "roe_pct": ("Return on equity", "%"), "nim_pct": ("Net interest margin", "%"),
    "cost_to_income_pct": ("Operating cost / income", "%"), "casa_pct": ("CASA share of deposits", "%"),
    "price_to_book": ("Indicative price-to-book multiple", "x"), "valuation_cr": ("Indicative valuation", "INR Cr"),
    "gnpa_pct": ("Gross NPA ratio", "%"), "nnpa_pct": ("Net NPA ratio", "%"), "pcr_pct": ("Provision coverage ratio", "%"),
    "car_pct": ("Capital adequacy ratio", "%"), "credit_deposit_ratio_pct": ("Credit-deposit ratio", "%"),
    "branches": ("Branch count", "#"), "customers_lakh": ("Customer base", "Lakh"), "employees": ("Headcount", "#"),
    "geo_overlap_pct": ("Share of target footprint overlapping acquirer's strong regions", "%"),
    "size_ratio": ("Target assets / acquirer assets", "ratio"),
    "app_mau_pct": ("Monthly active mobile-app users / customers", "%"),
    "digital_adoption_pct": ("Customers registered for digital channels", "%"),
    "digital_txn_share_pct": ("Share of transactions done digitally", "%"), "app_rating": ("App-store rating", "1-5"),
    "tech_modernity_score": ("Technology modernity (analyst score)", "1-5"),
    "core_banking_platform": ("Core banking system", "-"),
    "deal_id": ("Historical deal id", "-"), "deal_year": ("Year of announcement", "year"),
    "acquirer_assets_cr": ("Acquirer total assets at deal time", "INR Cr"),
    "synergy_achieved_pct": ("Synergies realised vs. announced after 36 months", "%"),
    "synergy_realized": ("1 = achieved >=80% of announced synergies (TARGET)", "0/1"),
    "customer_id": ("Unique customer id", "-"), "age": ("Customer age", "years"),
    "tenure_years": ("Relationship tenure", "years"), "avg_balance_inr": ("Average monthly balance", "INR"),
    "products_held": ("Number of products held", "#"), "has_loan": ("Holds a loan", "0/1"),
    "has_credit_card": ("Holds a credit card", "0/1"), "has_insurance": ("Holds an insurance policy", "0/1"),
    "salary_account": ("Salary / pension credited", "0/1"), "monthly_txn_count": ("Transactions per month", "#"),
    "login_frequency_30d": ("Digital log-ins in last 30 days", "#"),
    "branch_visits_30d": ("Branch visits in last 30 days", "#"), "digital_active": ("Active on digital (>=2 log-ins)", "0/1"),
    "complaints_12m": ("Complaints in last 12 months", "#"), "credit_score": ("Bureau score (may be blank)", "300-900"),
    "csat_score": ("Survey satisfaction (may be blank)", "1-10"),
    "est_annual_revenue_inr": ("Estimated annual revenue from customer", "INR"),
    "primary_branch_id": ("Home branch", "-"), "primary_product_id": ("Main product held", "-"),
    "fee_increase_flag": ("Customer faces higher fees after fee-schedule harmonisation", "0/1"),
    "branch_id": ("Unique branch id", "-"), "city": ("City", "-"), "state": ("State", "-"),
    "city_tier": ("City tier (1 = metro)", "1-3"), "branch_type": ("Branch format", "-"), "lat": ("Latitude", "deg"),
    "lon": ("Longitude", "deg"), "footfall_daily": ("Average daily footfall", "#"),
    "transactions_daily": ("Average daily counter transactions", "#"), "monthly_cost_lakh": ("Rent + staff + opex", "INR Lakh / month"),
    "staff_count": ("Branch staff", "#"), "distance_to_acq_branch_km": ("Distance to nearest acquirer branch", "km"),
    "product_id": ("Unique product id", "-"), "product_name": ("Product name", "-"), "product_type": ("Product family", "-"),
    "category": ("Deposit / Loan / Fee-based", "-"), "active_accounts": ("Active accounts", "#"),
    "annual_revenue_cr": ("Annual revenue", "INR Cr"), "revenue_per_account_inr": ("Revenue per account", "INR / yr"),
    "yoy_growth_pct": ("Year-on-year growth", "%"), "acquirer_equivalent_id": ("Closest acquirer product (blank = unique)", "-"),
    "feature_overlap_pct": ("Feature similarity to acquirer equivalent", "%"), "npa_rate_pct": ("Loan NPA rate", "%"),
    "merger_id": ("Completed merger id", "-"), "merger_year": ("Merger year", "year"),
    "target_bank_type": ("Type of the merged-in institution", "-"),
    "branch_closure_rate": ("Share of branches consolidated", "ratio"), "product_sunset_rate": ("Share of customers on sunset products", "ratio"),
    "fee_change_rate": ("Share of customers hit by fee change", "ratio"), "integration_severity": ("Integration difficulty index", "-0.3..0.5"),
    "branch_closure_flag": ("Home branch merged/closed", "0/1"), "product_sunset_flag": ("Main product sunset", "0/1"),
    "churned_12m": ("1 = left the bank within 12 months of merger (TARGET)", "0/1"),
    "metric": ("Client-bank metric name", "-"), "unit": ("Unit of the value", "-"), "description": ("Metric description", "-"),
    "attribute": ("Attribute name", "-"), "value": ("Attribute value", "-"),
}


def _style(ws):
    for c in ws[1]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = BODY_FONT
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(ws.iter_cols(min_row=1, max_row=min(ws.max_row, 200)), start=1):
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 3, 11), 48)


def write_workbook(path: Path, sheets: dict[str, pd.DataFrame], title: str, notes: list[str]) -> dict:
    path.parent.mkdir(parents=True, exist_ok=True)
    dict_rows = []
    for sheet, df in sheets.items():
        for col in df.columns:
            desc, unit = DICTIONARY.get(col, ("", ""))
            dict_rows.append((sheet, col, str(df[col].dtype), desc, unit, int(df[col].isna().sum())))
    dictionary = pd.DataFrame(dict_rows, columns=["sheet", "column", "dtype", "description", "unit", "null_count"])
    readme = pd.DataFrame({"OneRun AI - " + title: [
        f"Generated: {datetime.now():%d-%b-%Y %H:%M}",
        "All institutions, customers and figures are SYNTHETIC (fictional) and for demonstration only.",
        *notes,
        "Sheets: " + ", ".join(list(sheets) + ["Data_Dictionary"]),
        "Regenerate with:  python generate_data.py   (see README.md)",
    ]})
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        readme.to_excel(xw, sheet_name="README", index=False)
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name, index=False)
        dictionary.to_excel(xw, sheet_name="Data_Dictionary", index=False)
    wb = load_workbook(path)
    for ws in wb.worksheets:
        _style(ws)
    wb["README"].column_dimensions["A"].width = 110
    wb.save(path)
    return {name: (len(df), len(df.columns)) for name, df in sheets.items()}
